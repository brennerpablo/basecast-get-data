"""Parsers behind the raw-file previews, on local paths.

Spreadsheets and delimited text are shown positionally (column letters, the header is just row 1 or
wherever the publisher put it): the raw view shows what the source published and guesses nothing.
Parquet and JSON carry their own column names and types.
"""

from __future__ import annotations

import csv
import datetime as dt
import html.parser
import io
import math
import re
import warnings
import zipfile
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Any

import orjson
import polars as pl

from basecast_get_data.lake.kinds import kind_of
from basecast_get_data.schemas.common import Cell, ColumnType, GridColumn
from basecast_get_data.schemas.lake import JsonNode, SheetInfo, TextBlock, ZipMember


class NotTabular(Exception):
    """The file has no table to show in a grid."""


@dataclass
class Parsed:
    columns: list[GridColumn]
    frame: pl.DataFrame
    positional: bool
    delimiter: str | None = None
    tabular_path: str | None = None

    def rows(self, offset: int, limit: int) -> list[list[Cell]]:
        block = self.frame.slice(offset, limit)
        return [[to_cell(v, positional=self.positional) for v in row] for row in block.iter_rows()]


def letters(i: int) -> str:
    """0 → A, 25 → Z, 26 → AA: spreadsheet column names."""
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


MIDNIGHT = re.compile(r"^(\d{4}-\d{2}-\d{2})[ T]00:00:00$")


def to_cell(v: Any, *, positional: bool = False) -> Cell:
    if isinstance(v, str):
        # Mixed-type spreadsheet columns come back as text, dates included.
        if positional and (m := MIDNIGHT.match(v)):
            return m.group(1)
        return v
    if v is None or isinstance(v, bool):
        return v
    if isinstance(v, int):
        return v if abs(v) < 2**53 else str(v)
    if isinstance(v, float):
        if math.isnan(v) or math.isinf(v):
            return None
        # A spreadsheet shows 203, not 203.0.
        if positional and v.is_integer() and abs(v) < 1e15:
            return int(v)
        return v
    if isinstance(v, dt.datetime):
        if positional and v.time() == dt.time(0) and v.tzinfo is None:
            return v.date().isoformat()
        return v.isoformat()
    if isinstance(v, dt.date | dt.time):
        return v.isoformat()
    if isinstance(v, dt.timedelta):
        return str(v)
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, bytes | bytearray | memoryview):
        return f"<{len(v)} bytes>"
    return orjson.dumps(v, default=str).decode()


def column_type(dtype: pl.DataType) -> ColumnType:
    if dtype.is_numeric():
        return "number"
    if dtype.is_temporal():
        return "date"
    if dtype == pl.Boolean:
        return "boolean"
    if isinstance(dtype, pl.List | pl.Struct | pl.Array | pl.Object):
        return "json"
    return "text"


def _positional(frame: pl.DataFrame) -> Parsed:
    frame = frame.rename({c: letters(i) for i, c in enumerate(frame.columns)})
    return Parsed([GridColumn(name=c, type="text") for c in frame.columns], frame, positional=True)


# ---------- spreadsheets ----------


def sheet_names(path: Path) -> list[str]:
    import fastexcel

    return list(fastexcel.read_excel(str(path)).sheet_names)


def read_sheet(path: Path, sheet: str) -> Parsed:
    import fastexcel

    reader = fastexcel.read_excel(str(path))
    if sheet not in reader.sheet_names:
        raise KeyError(sheet)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        frame = reader.load_sheet(sheet, header_row=None).to_polars()
    return _positional(frame)


def sheet_infos(path: Path, parsed: dict[str, Parsed]) -> list[SheetInfo]:
    """Every sheet with its size; `parsed` holds the sheets already loaded (the rest load here)."""
    out = []
    for name in sheet_names(path):
        p = parsed.get(name) or read_sheet(path, name)
        parsed[name] = p
        out.append(SheetInfo(name=name, rows=p.frame.height, columns=p.frame.width))
    return out


# ---------- delimited text ----------

DELIMITERS = [",", "|", "\t", ";"]


def sniff_delimiter(sample: str) -> str | None:
    """The delimiter that splits the first lines into the same number of fields, if any."""
    lines = [ln for ln in sample.splitlines()[:40] if ln.strip()]
    if not lines:
        return None
    try:
        return csv.Sniffer().sniff("\n".join(lines[:20]), delimiters="".join(DELIMITERS)).delimiter
    except csv.Error:
        pass
    best, best_n = None, 0
    for d in DELIMITERS:
        counts = [ln.count(d) for ln in lines]
        n = min(counts)
        if n > best_n and max(counts) - n <= max(1, n // 10):
            best, best_n = d, n
    return best


def read_text(path: Path) -> Parsed:
    with path.open("rb") as f:
        sample = f.read(65536).decode("utf-8", errors="replace")
    delimiter = sniff_delimiter(sample)
    if delimiter is None:
        text = path.read_bytes().decode("utf-8", errors="replace")
        parsed = _positional(pl.DataFrame({"line": text.splitlines()}))
        return parsed
    # The widest row of the sample fixes the width: short rows get nulls, longer ones are cut.
    width = max((len(r) for r in csv.reader(io.StringIO(sample), delimiter=delimiter)), default=1)
    frame = pl.read_csv(
        path,
        has_header=False,
        separator=delimiter,
        schema={f"c{i}": pl.String for i in range(width)},
        truncate_ragged_lines=True,
        encoding="utf8-lossy",
    )
    parsed = _positional(frame)
    parsed.delimiter = delimiter
    return parsed


# ---------- parquet ----------


def read_parquet(path: Path) -> Parsed:
    frame = pl.read_parquet(path)
    cols = [GridColumn(name=c, type=column_type(t), source_type=str(t)) for c, t in frame.schema.items()]
    return Parsed(cols, frame, positional=False)


# ---------- JSON ----------


def _infer(values: list[Any]) -> tuple[ColumnType, list[Any]]:
    present = [v for v in values if v is not None]
    if present and all(isinstance(v, bool) for v in present):
        return "boolean", values
    if present and all(isinstance(v, int | float) and not isinstance(v, bool) for v in present):
        return "number", [float(v) if v is not None else None for v in values]
    if any(isinstance(v, dict | list) for v in present):
        return "json", [None if v is None else orjson.dumps(v).decode() for v in values]
    return "text", [None if v is None else str(v) for v in values]


def _frame(columns: dict[str, list[Any]]) -> tuple[list[GridColumn], pl.DataFrame]:
    out_cols, data = [], {}
    for name, values in columns.items():
        typ, conv = _infer(values)
        out_cols.append(GridColumn(name=name, type=typ))
        data[name] = pl.Series(name, conv, dtype=pl.Float64 if typ == "number" else None, strict=False)
    return out_cols, pl.DataFrame(data)


def _records(rows: list[dict[str, Any]]) -> dict[str, list[Any]]:
    keys: dict[str, None] = {}
    for r in rows:
        for k in r:
            keys.setdefault(str(k), None)
    return {k: [r.get(k) for r in rows] for k in keys}


def _parallel_arrays(obj: Any, path: str = "", depth: int = 0) -> tuple[str, dict[str, list]] | None:
    """The object with the most rows among those holding two or more arrays of the same length."""
    if not isinstance(obj, dict) or depth > 4:
        return None
    best: tuple[str, dict[str, list]] | None = None
    by_len: dict[int, dict[str, list]] = {}
    for k, v in obj.items():
        if isinstance(v, list) and len(v) >= 2 and all(not isinstance(x, dict | list) for x in v[:50]):
            by_len.setdefault(len(v), {})[str(k)] = v
    for n, group in by_len.items():
        if len(group) >= 2 and (best is None or n > len(next(iter(best[1].values())))):
            best = (path or "$", group)
    for k, v in obj.items():
        inner = _parallel_arrays(v, f"{path}.{k}" if path else str(k), depth + 1)
        if inner and (best is None or len(next(iter(inner[1].values()))) > len(next(iter(best[1].values())))):
            best = inner
    return best


def read_json(path: Path) -> Parsed:
    doc = orjson.loads(path.read_bytes())
    if isinstance(doc, dict) and doc.get("type") == "FeatureCollection":
        rows = []
        for f in doc.get("features") or []:
            props = dict(f.get("properties") or {})
            props["geometry"] = (f.get("geometry") or {}).get("type")
            rows.append(props)
        cols, frame = _frame(_records(rows))
        return Parsed(cols, frame, positional=False, tabular_path="features[].properties")
    if isinstance(doc, list) and doc and all(isinstance(r, dict) for r in doc[:100]):
        cols, frame = _frame(_records([r for r in doc if isinstance(r, dict)]))
        return Parsed(cols, frame, positional=False, tabular_path="$[]")
    found = _parallel_arrays(doc)
    if found:
        path_label, group = found
        cols, frame = _frame(group)
        return Parsed(cols, frame, positional=False, tabular_path=path_label)
    raise NotTabular("no array of records or parallel arrays in this JSON")


def json_tree(value: Any, key: str | None = None, depth: int = 0) -> JsonNode:
    """The document as a tree, cut down: arrays keep their first items, objects their first keys."""
    if isinstance(value, dict):
        items = list(value.items())
        children = [json_tree(v, str(k), depth + 1) for k, v in items[:200]] if depth < 8 else []
        return JsonNode(key=key, type="object", length=len(items), children=children)
    if isinstance(value, list):
        children = [json_tree(v, str(i), depth + 1) for i, v in enumerate(value[:5])] if depth < 8 else []
        return JsonNode(key=key, type="array", length=len(value), children=children)
    if isinstance(value, bool):
        return JsonNode(key=key, type="boolean", value=str(value).lower())
    if isinstance(value, int | float):
        return JsonNode(key=key, type="number", value=repr(value))
    if value is None:
        return JsonNode(key=key, type="null")
    text = str(value)
    return JsonNode(key=key, type="string", value=text if len(text) <= 500 else text[:500] + "…")


def read_json_tree(path: Path) -> JsonNode:
    return json_tree(orjson.loads(path.read_bytes()))


# ---------- zip ----------


def zip_members(path: Path) -> list[ZipMember]:
    with zipfile.ZipFile(path) as z:
        return [
            ZipMember(
                name=i.filename,
                kind=kind_of(i.filename),  # type: ignore[arg-type]
                bytes=i.file_size,
                compressed_bytes=i.compress_size,
                modified=dt.datetime(*i.date_time).isoformat() if i.date_time[0] >= 1980 else None,
            )
            for i in z.infolist()
            if not i.is_dir()
        ]


def extract_member(path: Path, member: str, dest: Path) -> Path:
    with zipfile.ZipFile(path) as z:
        if member not in z.namelist():
            raise KeyError(member)
        dest.parent.mkdir(parents=True, exist_ok=True)
        tmp = dest.with_suffix(dest.suffix + ".part")
        with z.open(member) as src, tmp.open("wb") as out:
            while chunk := src.read(1 << 20):
                out.write(chunk)
        tmp.replace(dest)
    return dest


# ---------- text documents ----------


def _pptx_text(text: str) -> str:
    # PowerPoint stores soft line breaks as vertical tabs.
    return re.sub(r"\n{3,}", "\n\n", text.replace("\x0b", "\n")).strip()


def pptx_blocks(path: Path) -> list[TextBlock]:
    from pptx import Presentation

    blocks = []
    for i, slide in enumerate(Presentation(str(path)).slides, start=1):
        title_shape = slide.shapes.title
        title = _pptx_text(title_shape.text) if title_shape is not None and title_shape.has_text_frame else ""
        title = " ".join(title.split())
        parts = []
        for shape in slide.shapes:
            if shape is title_shape:
                continue
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(_pptx_text(shape.text_frame.text))
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append("\t".join(c.text.strip() for c in row.cells))
        blocks.append(TextBlock(title=f"Slide {i}" + (f": {title}" if title else ""), text="\n".join(parts)))
    return blocks


def docx_blocks(path: Path) -> list[TextBlock]:
    from docx import Document

    doc = Document(str(path))
    blocks: list[TextBlock] = []
    title: str | None = None
    lines: list[str] = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if (p.style.name or "").lower().startswith(("heading", "title")):
            if title or lines:
                blocks.append(TextBlock(title=title, text="\n".join(lines)))
            title, lines = text, []
        else:
            lines.append(text)
    if title or lines:
        blocks.append(TextBlock(title=title, text="\n".join(lines)))
    for n, table in enumerate(doc.tables, start=1):
        rows = ["\t".join(c.text.strip() for c in r.cells) for r in table.rows]
        blocks.append(TextBlock(title=f"Table {n}", text="\n".join(rows)))
    return blocks


class _HtmlText(html.parser.HTMLParser):
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "table"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.out: list[str] = []
        self.skip = 0
        self.title = ""
        self._in_title = False

    def handle_starttag(self, tag: str, attrs: Any) -> None:
        if tag in ("script", "style", "noscript"):
            self.skip += 1
        if tag == "title":
            self._in_title = True
        if tag in self.BLOCK:
            self.out.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in ("script", "style", "noscript") and self.skip:
            self.skip -= 1
        if tag == "title":
            self._in_title = False
        if tag in ("td", "th"):
            self.out.append("\t")

    def handle_data(self, data: str) -> None:
        if self._in_title:
            self.title += data
        elif not self.skip:
            self.out.append(data)


def html_blocks(path: Path) -> list[TextBlock]:
    parser = _HtmlText()
    parser.feed(path.read_bytes().decode("utf-8", errors="replace"))
    text = re.sub(r"[ \t\r\f\v]+", " ", "".join(parser.out))
    text = re.sub(r"\n\s*\n+", "\n\n", text).strip()
    return [TextBlock(title=parser.title.strip() or None, text=text)]


def text_blocks(path: Path, kind: str) -> list[TextBlock]:
    if kind == "slides":
        return pptx_blocks(path)
    if kind == "document":
        return docx_blocks(path)
    if kind == "html":
        return html_blocks(path)
    raise NotTabular(f"no text view for {kind}")


def read_tabular(path: Path, kind: str, sheet: str | None = None) -> Parsed:
    if kind == "sheet":
        return read_sheet(path, sheet or sheet_names(path)[0])
    if kind == "text":
        return read_text(path)
    if kind == "parquet":
        return read_parquet(path)
    if kind in ("json", "geojson"):
        return read_json(path)
    raise NotTabular(f"no grid view for {kind}")
